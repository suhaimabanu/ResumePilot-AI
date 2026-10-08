"""Rebuild a DOCX for PDF-sourced resumes, mirroring the original's measured layout.

A PDF has no editable template, so we reproduce what we measured from it: page size, margins,
per-paragraph font family/size/weight/colour/alignment, bullets, right-aligned dates (tabs),
vertical spacing and the profile photo. Multi-column PDFs are linearised (documented limitation).
"""
from __future__ import annotations

import io
from pathlib import Path

import docx
import pymupdf as fitz
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..core.models import Block, ResumeSource
from ..parsing.sections import strip_bullet

ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}
FONT_MAP = {"ArialMT": "Arial", "Helvetica": "Arial", "TimesNewRomanPSMT": "Times New Roman", "Times": "Times New Roman",
            "Times-Roman": "Times New Roman", "CalibriLight": "Calibri Light", "Calibri-Light": "Calibri Light"}


def extract_profile_photo(pdf_path: Path) -> tuple[bytes, float] | None:
    """Largest roughly-square image on page 1 -> (png bytes, width in inches)."""
    doc = fitz.open(str(pdf_path))
    try:
        if not doc.page_count:
            return None
        page = doc[0]
        best = None
        for info in page.get_image_info(xrefs=True):
            x0, y0, x1, y1 = info["bbox"]
            w, h = x1 - x0, y1 - y0
            if w < 40 or h < 40 or not (0.6 <= w / h <= 1.6) or not info.get("xref"):
                continue
            if best is None or w * h > best[1] * best[2]:
                best = (info["xref"], w, h)
        if not best:
            return None
        pix = fitz.Pixmap(doc, best[0])
        if pix.alpha or pix.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        return pix.tobytes("png"), best[1] / 72
    finally:
        doc.close()


def _font_name(name: str | None) -> str | None:
    if not name:
        return None
    return FONT_MAP.get(name, name)


def _style_paragraph(p, block: Block, text_width_in: float, body_size: float) -> None:
    st = block.style
    fmt = p.paragraph_format
    fmt.space_after = Pt(0)
    gap = st.space_before or 0
    size = st.size or body_size
    fmt.space_before = Pt(max(0, min(14, gap - size * 0.3))) if gap else Pt(0)
    if st.align in ALIGN:
        p.alignment = ALIGN[st.align]
    if "\t" in block.text:
        fmt.tab_stops.add_tab_stop(Inches(text_width_in), WD_TAB_ALIGNMENT.RIGHT)


def _add_runs(p, text: str, block: Block, body_size: float) -> None:
    st = block.style
    run = p.add_run(text)
    run.bold = st.bold or block.kind == "heading"
    run.italic = st.italic
    run.font.size = Pt(st.size or body_size)
    font = _font_name(st.font)
    if font:
        run.font.name = font
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font)
    if st.color and st.color != "000000":
        try:
            run.font.color.rgb = RGBColor.from_string(st.color)
        except ValueError:
            pass


def build_docx_from_blocks(resume: ResumeSource, texts: dict[str, str], pdf_source: Path, out_path: Path,
                           spacing_factor: float = 1.0) -> dict:
    fp = resume.fingerprint
    document = docx.Document()
    section = document.sections[0]
    if fp.page_size == "A4":
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    m = fp.margins or {}
    section.left_margin = Inches(max(0.3, min(1.25, m.get("left", 0.75))))
    section.right_margin = Inches(max(0.3, min(1.25, m.get("right", 0.75))))
    section.top_margin = Inches(max(0.3, min(1.25, m.get("top", 0.6))))
    section.bottom_margin = Inches(max(0.3, min(1.0, m.get("bottom", 0.6))))
    text_width = (section.page_width - section.left_margin - section.right_margin) / 914400

    sizes = [b.style.size for b in resume.blocks if b.style.size and b.kind in ("paragraph", "bullet")]
    body_size = sorted(sizes)[len(sizes) // 2] if sizes else 10.5
    normal = document.styles["Normal"]
    normal.font.size = Pt(body_size)
    common_font = _font_name(fp.fonts[0]) if fp.fonts else None
    fonts = [b.style.font for b in resume.blocks if b.style.font]
    if fonts:
        common_font = _font_name(max(set(fonts), key=fonts.count))
    if common_font:
        normal.font.name = common_font
    normal.paragraph_format.space_after = Pt(0)

    photo = extract_profile_photo(pdf_source) if fp.profile_photo else None
    blocks = [b for b in resume.blocks if b.kind != "empty"]
    first_section = next((i for i, b in enumerate(blocks) if b.role == "section_heading"), len(blocks))
    header_blocks, body_blocks = blocks[:first_section], blocks[first_section:]

    container = document
    if photo and header_blocks:
        table = document.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = table.rows[0].cells
        photo_w = max(0.8, min(1.6, photo[1]))
        left.width, right.width = Inches(text_width - photo_w - 0.1), Inches(photo_w + 0.1)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.add_run().add_picture(io.BytesIO(photo[0]), width=Inches(photo_w))
        container = left
        first_par = left.paragraphs[0]
    else:
        first_par = None

    def emit(target, block: Block, reuse=None):
        text = texts.get(block.id, block.text)
        if block.kind == "bullet":
            _, body = strip_bullet(text)
            p = reuse or target.add_paragraph(style="List Bullet" if target is document else None)
            if target is not document:
                body = "• " + body
            _style_paragraph(p, block, text_width, body_size)
            _add_runs(p, body, block, body_size)
        else:
            p = reuse or target.add_paragraph()
            _style_paragraph(p, block, text_width, body_size)
            _add_runs(p, text, block, body_size)
        if spacing_factor < 1.0 and p.paragraph_format.space_before:
            p.paragraph_format.space_before = Pt(p.paragraph_format.space_before.pt * spacing_factor)

    for i, b in enumerate(header_blocks):
        emit(container, b, reuse=first_par if (i == 0 and first_par is not None) else None)
    for b in body_blocks:
        emit(document, b)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return {"rebuilt_from_pdf": True, "photo": bool(photo), "spacing_factor": spacing_factor}

