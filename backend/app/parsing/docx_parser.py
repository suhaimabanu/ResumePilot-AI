"""Parse an uploaded DOCX into raw paragraphs + a template fingerprint."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from lxml import etree

from ..core.models import BlockStyle, TemplateFingerprint
from .docx_xml import (A_NS, NSMAP, W, count_not_in_fallback, enumerate_paragraphs, header_footer_parts,
                       iter_parts_xml, own_runs, paragraph_text)
from .sections import RawParagraph

EMU_PER_INCH = 914400


class _StyleInfo:
    def __init__(self, doc: DocxDocument) -> None:
        self.by_id: dict[str, etree._Element] = {}
        styles_el = doc.styles.element
        for s in styles_el.findall(W + "style"):
            sid = s.get(W + "styleId")
            if sid:
                self.by_id[sid] = s
        self.default_para = next((sid for sid, s in self.by_id.items()
                                  if s.get(W + "type") == "paragraph" and s.get(W + "default") in ("1", "true")), None)

    def chain(self, style_id: str | None):
        seen = set()
        while style_id and style_id in self.by_id and style_id not in seen:
            seen.add(style_id)
            el = self.by_id[style_id]
            yield el
            based = el.find(W + "basedOn")
            style_id = based.get(W + "val") if based is not None else None

    def name(self, style_id: str | None) -> str:
        if style_id and style_id in self.by_id:
            n = self.by_id[style_id].find(W + "name")
            return n.get(W + "val") if n is not None else style_id
        return ""

    def has_numbering(self, style_id: str | None) -> bool:
        return any(el.find(f"{W}pPr/{W}numPr") is not None for el in self.chain(style_id))

    def rpr_value(self, style_id: str | None, tag: str) -> etree._Element | None:
        for el in self.chain(style_id):
            found = el.find(f"{W}rPr/{W}{tag}")
            if found is not None:
                return found
        return None


def _on(el: etree._Element | None) -> bool | None:
    if el is None:
        return None
    return el.get(W + "val") not in ("0", "false", "off")


def _run_prop(run: etree._Element, styles: _StyleInfo, p_style: str | None, tag: str):
    rpr = run.find(W + "rPr")
    if rpr is not None:
        el = rpr.find(W + tag)
        if el is not None:
            return el
        rstyle = rpr.find(W + "rStyle")
        if rstyle is not None:
            el = styles.rpr_value(rstyle.get(W + "val"), tag)
            if el is not None:
                return el
    return styles.rpr_value(p_style or styles.default_para, tag)


def _paragraph_raw(handle, styles: _StyleInfo) -> RawParagraph:
    p = handle.element
    text = paragraph_text(p)
    ppr = p.find(W + "pPr")
    p_style = None
    if ppr is not None and (ps := ppr.find(W + "pStyle")) is not None:
        p_style = ps.get(W + "val")
    style_name = styles.name(p_style).lower()
    is_list = (ppr is not None and ppr.find(W + "numPr") is not None) or styles.has_numbering(p_style) or "list" in style_name
    is_heading = bool(re.match(r"(heading\s*\d|title)", style_name))

    runs = [r for r in own_runs(p) if any(t.text and t.text.strip() for t in r.findall(W + "t"))]
    bold_flags = [bool(_on(_run_prop(r, styles, p_style, "b"))) for r in runs]
    caps_flags = [bool(_on(_run_prop(r, styles, p_style, "caps"))) for r in runs]
    letters = re.sub(r"[^A-Za-z]", "", text)
    all_caps = (len(letters) >= 3 and letters.isupper()) or (bool(caps_flags) and all(caps_flags))
    style = BlockStyle()
    if runs:
        sz = _run_prop(runs[0], styles, p_style, "sz")
        if sz is not None and sz.get(W + "val", "").isdigit():
            style.size = int(sz.get(W + "val")) / 2
        fonts = _run_prop(runs[0], styles, p_style, "rFonts")
        if fonts is not None:
            style.font = fonts.get(W + "ascii") or fonts.get(W + "hAnsi")
        color = _run_prop(runs[0], styles, p_style, "color")
        if color is not None and color.get(W + "val") not in (None, "auto"):
            style.color = color.get(W + "val")
        style.bold = bool(bold_flags) and all(bold_flags)
    if ppr is not None and (jc := ppr.find(W + "jc")) is not None:
        style.align = {"both": "justify"}.get(jc.get(W + "val"), jc.get(W + "val"))
    return RawParagraph(
        text=text,
        location=handle.location,
        is_list=is_list,
        is_heading_style=is_heading,
        all_bold=bool(bold_flags) and all(bold_flags),
        all_caps=all_caps,
        style=style,
    )


def parse_docx(path: Path) -> tuple[list[RawParagraph], TemplateFingerprint, int | None]:
    doc = docx.Document(str(path))
    styles = _StyleInfo(doc)
    raw = [_paragraph_raw(h, styles) for h in enumerate_paragraphs(doc)]
    fingerprint = docx_fingerprint(doc)
    fingerprint.paragraphs = len(raw)
    return raw, fingerprint, docx_app_pages(path)


def docx_app_pages(path: Path) -> int | None:
    """Page count Word stored in docProps/app.xml (a hint only — rendering is authoritative)."""
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("docProps/app.xml").decode("utf-8", "ignore")
        m = re.search(r"<Pages>(\d+)</Pages>", xml)
        return int(m.group(1)) if m else None
    except Exception:
        return None


def _page_size_name(w_twips: int, h_twips: int) -> str:
    w_in, h_in = w_twips / 1440, h_twips / 1440
    for name, (a, b) in {"Letter": (8.5, 11), "A4": (8.27, 11.69), "Legal": (8.5, 14)}.items():
        if abs(min(w_in, h_in) - a) < 0.08 and abs(max(w_in, h_in) - b) < 0.08:
            return name
    return f"{w_in:.2f}x{h_in:.2f}in"


def docx_fingerprint(doc: DocxDocument) -> TemplateFingerprint:
    fp = TemplateFingerprint()
    body = doc.element.body
    sect = body.findall(f".//{W}sectPr")
    fp.section_breaks = max(0, len(sect) - 1)
    if sect:
        last = sect[-1]
        pgsz = last.find(W + "pgSz")
        if pgsz is not None and pgsz.get(W + "w"):
            fp.page_size = _page_size_name(int(pgsz.get(W + "w")), int(pgsz.get(W + "h")))
        pgmar = last.find(W + "pgMar")
        if pgmar is not None:
            for side in ("top", "bottom", "left", "right"):
                v = pgmar.get(W + side)
                if v and v.lstrip("-").isdigit():
                    fp.margins[side] = round(int(v) / 1440, 2)
        cols = [int(c.get(W + "num")) for s in sect for c in s.findall(W + "cols") if (c.get(W + "num") or "").isdigit()]
        fp.columns = max(cols) if cols else 1

    fonts: set[str] = set()
    sizes: set[float] = set()
    colors: set[str] = set()
    images = text_boxes = hyperlinks = tables = 0
    profile_photo = False
    used_styles = {el.get(W + "val") for root in iter_parts_xml(doc) for tag in ("pStyle", "rStyle", "tblStyle")
                   for el in root.iter(W + tag)}
    styles_info = _StyleInfo(doc)
    if styles_info.default_para:
        used_styles.add(styles_info.default_para)
    style_roots = [el for sid in used_styles for el in styles_info.chain(sid)]
    doc_defaults = doc.styles.element.find(W + "docDefaults")
    if doc_defaults is not None:
        style_roots.append(doc_defaults)
    for root in list(iter_parts_xml(doc)) + style_roots:
        for f in root.iter(W + "rFonts"):
            for attr in ("ascii", "hAnsi"):
                if f.get(W + attr):
                    fonts.add(f.get(W + attr))
        for s in root.iter(W + "sz"):
            v = s.get(W + "val")
            if v and v.isdigit():
                sizes.add(int(v) / 2)
        for c in root.iter(W + "color"):
            v = c.get(W + "val")
            if v and v != "auto":
                colors.add(v.upper())
        for shd in root.iter(W + "shd"):
            v = shd.get(W + "fill")
            if v and v not in ("auto", "FFFFFF"):
                colors.add(v.upper())
        if root in style_roots:
            continue
        images += count_not_in_fallback(root, ".//a:blip") + count_not_in_fallback(root, ".//v:imagedata")
        text_boxes += count_not_in_fallback(root, ".//w:txbxContent")
        hyperlinks += len(root.findall(f".//{W}hyperlink"))
        tables += count_not_in_fallback(root, ".//w:tbl")
        for ext in etree.XPath(".//wp:inline/wp:extent | .//wp:anchor/wp:extent", namespaces=NSMAP)(root):
            try:
                cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            except (TypeError, ValueError):
                continue
            if cy and 0.6 <= cx / cy <= 1.6 and min(cx, cy) >= 0.6 * EMU_PER_INCH:
                # A roughly square picture of at least ~0.6" is very likely a profile photo.
                container = ext.getparent()
                if container is not None and container.find(f".//{{{A_NS}}}blip") is not None:
                    profile_photo = True
    hf = header_footer_parts(doc)
    fp.headers = sum(1 for k, p in hf if k == "header" and _has_content(p.element))
    fp.footers = sum(1 for k, p in hf if k == "footer" and _has_content(p.element))
    fp.fonts = sorted(fonts)
    fp.font_sizes = sorted(sizes)
    fp.colors = sorted(colors)
    fp.images, fp.text_boxes, fp.hyperlinks, fp.tables = images, text_boxes, hyperlinks, tables
    fp.profile_photo = profile_photo
    fp.styles = sorted({ps.get(W + "val") for ps in body.iter(W + "pStyle") if ps.get(W + "val")})
    return fp


def _has_content(root: etree._Element) -> bool:
    return any((t.text or "").strip() for t in root.iter(W + "t")) or root.find(f".//{{{A_NS}}}blip") is not None
