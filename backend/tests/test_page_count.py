"""Page-count detection and overflow optimization (never shrinks fonts)."""
from __future__ import annotations

import shutil

import docx
from docx.shared import Pt

from app.ats.evidence import assess_keywords
from app.ats.jd_parser import parse_jd_text
from app.core.jobs import Job, JobContext
from app.core.models import Change
from app.core.store import store
from app.documents.converter import docx_to_pdf
from app.parsing.pdf_parser import pdf_page_count
from app.tailoring import pipeline

from .conftest import needs_converter

BULLET = "Delivered internal reporting tools with Python and SQL for the operations and finance teams."


def _make(path, n):
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(11)
    d.add_paragraph("Morgan Lee").runs[0].bold = True
    d.add_paragraph("morgan.lee@example.org | +44 20 7946 0000")
    d.add_paragraph("Summary").runs[0].bold = True
    d.add_paragraph("Data engineer who builds reporting tools with Python and SQL for operations teams.")
    d.add_paragraph("Experience").runs[0].bold = True
    d.add_paragraph("Data Engineer | Riverbank Group\tJan 2019 – Present")
    for _ in range(n):
        p = d.add_paragraph(BULLET, style="List Bullet")
        p.paragraph_format.space_after = Pt(9)
    d.add_paragraph("Education").runs[0].bold = True
    d.add_paragraph("BSc Mathematics | University of Leeds\t2015 – 2018")
    d.save(path)


@needs_converter
def test_overflow_is_detected_and_optimized(tmp_path, no_ai):
    # Find the largest bullet count that still fits on one page.
    path = tmp_path / "full.docx"
    def pages(count: int) -> int:
        _make(path, count)
        return pdf_page_count(docx_to_pdf(path, tmp_path)[0])

    n = 16
    while pages(n) > 1:
        n -= 2
    while pages(n + 1) == 1:
        n += 1
    _make(path, n)

    s = store.create()
    dst = s.directory / "resume_source.docx"
    shutil.copy(path, dst)
    s.resume_path, s.resume_filename = dst, "full.docx"
    pipeline.run_upload(s, JobContext(Job("j", "s", "upload", [])), "docx")
    assert s.resume.page_count == 1
    s.jd_text = "Data Engineer\nRequirements:\n- Python and SQL reporting for operations and finance teams\n- 3+ years"
    s.jd_analysis = parse_jd_text(s.jd_text)
    s.keyword_evidence = assess_keywords(s.jd_analysis, s.resume.full_text)

    bullets = [b for b in s.resume.blocks if b.role == "experience_bullet"]
    longer = BULLET[:-1] + ", delivered with Python and SQL reporting tools for the operations and finance teams."
    s.changes = [Change(change_id=f"CHG{i + 1:03d}", block_id=b.id, section=b.section, original=b.text,
                        suggested=longer, status="accepted") for i, b in enumerate(bullets[:3])]
    ctx = JobContext(Job("j", "s", "generation", []))
    result = pipeline.run_generation(s, ctx)
    pages = result["pages"]
    assert pages["original"] == 1
    assert pages["optimization"], "overflow should have triggered layout optimization"
    assert pages["tailored"] == 1, pages
    # Fonts are never shrunk to force a fit.
    sizes = {r.font.size.pt for p in docx.Document(str(s.files[result["downloads"]["docx"]["file_id"]].path)).paragraphs
             for r in p.runs if r.font.size}
    assert not sizes or min(sizes) >= 11 or min(sizes) == 11
    store.delete(s.session_id)
