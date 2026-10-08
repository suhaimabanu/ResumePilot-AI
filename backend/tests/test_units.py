"""Unit tests: upload validation, parsing, matching, ATS scoring, fact checking, DOCX editing,
completeness, AI response handling."""
from __future__ import annotations

import io
import zipfile

import docx
import pytest

from app.ai import service as ai_service
from app.ai.base import AIResponseError, parse_json_text
from app.ats.evidence import assess_keywords, verify_quote
from app.ats.jd_parser import parse_jd_text
from app.ats.matching import match_keyword
from app.ats.scoring import compare_scores, score_resume
from app.core.files import validate_upload
from app.core.jobs import UserFacingError
from app.parsing.docx_xml import W, enumerate_paragraphs, own_runs, paragraph_text, set_paragraph_text
from app.parsing.resume_builder import parse_resume
from app.validation.content import validate_identity
from app.validation.fact_check import check_change
from app.validation.resume_completeness_validator import validate_completeness

from .fake_ai import FakeProvider


# ---------------------------------------------------------------- upload validation
def test_upload_rejects_bad_files(sample_docx):
    data = sample_docx.read_bytes()
    assert validate_upload("cv.docx", None, data, {"docx", "pdf"}) == "docx"
    with pytest.raises(UserFacingError):
        validate_upload("cv.exe", None, b"MZ\x90\x00binary", {"docx", "pdf"})
    with pytest.raises(UserFacingError):
        validate_upload("cv.docx", None, b"not a zip at all", {"docx", "pdf"})  # signature mismatch
    with pytest.raises(UserFacingError):
        validate_upload("cv.pdf", None, b"%PDF-1.4 garbage", {"docx", "pdf"})  # corrupt pdf
    with pytest.raises(UserFacingError):
        validate_upload("cv.txt", None, b"hello world", {"docx", "pdf"})  # unsupported for resume
    with pytest.raises(UserFacingError):
        validate_upload("cv.docx", None, b"", {"docx", "pdf"})
    big = b"%PDF-" + b"0" * (11 * 1024 * 1024)
    with pytest.raises(UserFacingError, match="too large"):
        validate_upload("cv.pdf", None, big, {"docx", "pdf"})
    # A zip that is not a Word document
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "x")
    with pytest.raises(UserFacingError):
        validate_upload("cv.docx", None, buf.getvalue(), {"docx", "pdf"})


# ---------------------------------------------------------------- parsing
def test_docx_parsing_uses_actual_resume(sample_docx):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    assert r.name == "Jordan Avery"
    assert r.contact.email == "jordan.avery@example.com"
    assert len(r.experience) == 3
    assert [e.organization for e in r.experience] == ["Northwind Analytics", "Bluebird Logistics", "Copperleaf Studio"]
    assert r.experience[0].dates.startswith("Mar 2021")
    assert len(r.projects) == 2 and len(r.education) == 1
    assert "AWS (EC2, S3, Lambda)" in r.skills
    assert r.fingerprint.profile_photo and r.fingerprint.tables == 1 and r.fingerprint.footers == 1
    assert {s.key for s in r.sections} >= {"summary", "skills", "experience", "projects", "education", "certifications"}
    # locked vs editable
    roles = {b.role: b.editable for b in r.blocks}
    assert roles["experience_header"] is False and roles["experience_bullet"] is True
    assert roles["name"] is False and roles["contact"] is False
    assert r.claims and all(c.id.startswith("SRC") for c in r.claims)


def test_unreadable_resume_stops(tmp_path):
    d = docx.Document()
    d.add_paragraph("Hi")
    p = tmp_path / "empty.docx"
    d.save(p)
    with pytest.raises(UserFacingError, match="couldn't reliably read"):
        parse_resume(p, "docx", "empty.docx")


def test_jd_parsing(sample_jd):
    jd = parse_jd_text(sample_jd)
    terms = {k.term for k in jd.keywords}
    assert {"Python", "FastAPI", "Kubernetes", "Terraform", "AWS", "Docker"} <= terms
    assert jd.years_experience == 5
    assert jd.job_title.startswith("Senior Python Backend Engineer")
    assert any("Bachelor" in e for e in jd.education_requirements)
    imp = {k.term: k.importance for k in jd.keywords}
    assert imp["Kubernetes"] == "required" and imp["Kafka"] == "preferred"


# ---------------------------------------------------------------- matching
@pytest.mark.parametrize("term,text,expected", [
    ("REST API", "Built RESTful APIs", "alias"),
    ("PostgreSQL", "Postgres, Redis", "alias"),
    ("AWS", "Amazon Web Services (EC2)", "alias"),
    ("Python", "Python developer", "exact"),
    ("REST API", "REST APIs for clients", "exact"),
    ("Agile", "worked in two-week sprints", "related"),
])
def test_keyword_matching_positive(term, text, expected):
    assert match_keyword(term, text)[0] == expected


@pytest.mark.parametrize("term,text", [
    ("React", "Django developer"), ("Java", "Python and JavaScript"), ("Azure", "AWS and GCP"),
    ("Go", "go to market strategy"), ("C", "C++ and C# developer"), ("Kubernetes", "Docker containers"),
])
def test_keyword_matching_negative(term, text):
    assert match_keyword(term, text)[0] == "none"


# ---------------------------------------------------------------- scoring
def test_ats_score_is_deterministic_and_explained(sample_docx, sample_jd):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    jd = parse_jd_text(sample_jd)
    a, b = score_resume(r, jd), score_resume(r, jd)
    assert a == b
    assert [c["max"] for c in a["categories"]] == [30, 20, 15, 15, 10, 10]
    assert 0 <= a["total"] <= 100
    kw = a["categories"][0]
    assert "Kubernetes" in kw["missing"] and "Python" in kw["matched"]
    assert all(c["explanation"] for c in a["categories"])
    texts = {bl.id: bl.text for bl in r.blocks}
    bid = next(bl.id for bl in r.blocks if bl.text.startswith("Developed RESTful services"))
    texts[bid] = "Developed REST APIs in Python and FastAPI for the customer analytics platform."
    after = score_resume(r, jd, texts)
    cmp = compare_scores(a, after)
    assert cmp["keyword_coverage"]["tailored"] >= cmp["keyword_coverage"]["original"]


def test_keyword_stuffing_is_penalized(sample_docx, sample_jd):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    jd = parse_jd_text(sample_jd)
    texts = {b.id: b.text for b in r.blocks}
    bid = next(b.id for b in r.blocks if b.role == "summary")
    texts[bid] = texts[bid] + " Python " * 12
    stuffed = score_resume(r, jd, texts)
    assert "Python" in stuffed["categories"][0]["stuffing"]


def test_evidence_never_comes_from_jd(sample_docx, sample_jd):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    jd = parse_jd_text(sample_jd)
    ev = {e.term: e for e in assess_keywords(jd, r.full_text, {"Kubernetes": "Deploy and operate services on Amazon Web Services using Docker and Kubernetes"})}
    assert ev["Kubernetes"].status == "missing" and not ev["Kubernetes"].usable
    assert ev["Python"].status == "supported" and ev["Python"].usable
    assert verify_quote("adding Redis caching", r.full_text)
    assert not verify_quote("led Kubernetes migration", r.full_text)


# ---------------------------------------------------------------- anti-fabrication
@pytest.fixture
def parsed(sample_docx, sample_jd):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    jd = parse_jd_text(sample_jd)
    return r, jd, assess_keywords(jd, r.full_text)


def _block(r, prefix):
    return next(b for b in r.blocks if b.text.startswith(prefix))


def test_fact_check_blocks_fabrications(parsed):
    r, jd, ev = parsed
    src = r.full_text
    b = _block(r, "Containerised 12 services")
    assert not check_change(b, "Containerised 12 services with Docker and Kubernetes on AWS ECS.", src, ev).passed
    b = _block(r, "Designed PostgreSQL")
    v = check_change(b, "Designed PostgreSQL schemas, improving query speed by 50%.", src, ev)
    assert not v.passed and any("50" in i for i in v.issues)
    b = _block(r, "Reduced average API response time")
    v = check_change(b, "Reduced API response time using Redis caching and query optimisation.", src, ev)
    assert not v.passed and any("35%" in i for i in v.issues)  # dropped metric
    b = _block(r, "Built backend applications")
    assert not check_change(b, "Built Django applications for Microsoft and 200+ clients.", src, ev).passed
    b = next(x for x in r.blocks if x.role == "summary")
    assert not check_change(b, "Software engineer with 10 years of experience in Python.", src, ev).passed


def test_fact_check_accepts_supported_rewrites(parsed):
    r, jd, ev = parsed
    src = r.full_text
    b = _block(r, "Developed RESTful services")
    assert check_change(b, "Developed Python REST APIs with FastAPI for the customer analytics platform.", src, ev).passed
    b = _block(r, "Cloud & DevOps:")
    v = check_change(b, "Cloud & DevOps: Docker, AWS (EC2, S3, Lambda), GitHub Actions, Linux", src, ev)
    assert v.passed, v.issues
    v = check_change(b, "Cloud & DevOps: Docker, Kubernetes, AWS (EC2, S3, Lambda), GitHub Actions, Linux", src, ev)
    assert not v.passed
    v = check_change(b, "Cloud & DevOps: Docker, AWS (EC2, S3, Lambda)", src, ev)
    assert not v.passed and any("Removes skills" in i for i in v.issues)


def test_gemini_model_fallback_chain(monkeypatch):
    from app.ai import gemini
    p = gemini.GeminiProvider(api_key="k", model="m-main", fallback_models=["m-2", "m-lite", "m-main"])
    assert p.models == ["m-main", "m-2", "m-lite"]  # de-duplicated, primary first
    assert p._candidate_models("structure") == ["m-main", "m-2", "m-lite"]
    assert p._candidate_models("tailor") == ["m-main", "m-2"]  # lite never rewrites a resume
    monkeypatch.setitem(gemini._cooldown, "m-main", 1e12)  # overloaded main model is skipped while cooling down
    assert p._candidate_models("tailor") == ["m-2"]


def test_fact_check_blocks_seniority_and_dropped_facts(parsed):
    r, jd, ev = parsed
    src = r.full_text
    summary = next(x for x in r.blocks if x.role == "summary")
    assert "principal" not in src.lower()
    v = check_change(summary, "Principal " + summary.text, src, ev)
    assert not v.passed and any("seniority" in i for i in v.issues)
    from app.validation.fact_check import SENIORITY_RE
    assert not SENIORITY_RE.search("Helped lead the migration of services")  # the verb "lead" is not a rank
    assert SENIORITY_RE.search("Lead Backend Engineer") and SENIORITY_RE.search("a senior python developer")
    b = _block(r, "Developed RESTful services")
    v = check_change(b, "Developed REST APIs for the platform.", src, ev)  # FastAPI dropped
    assert not v.passed and any("Drops 'FastAPI'" in i for i in v.issues), v.issues


# ---------------------------------------------------------------- DOCX template editing
def test_set_paragraph_text_preserves_run_formatting(sample_docx, tmp_path):
    d = docx.Document(str(sample_docx))
    handles = enumerate_paragraphs(d)
    skills = next(h for h in handles if paragraph_text(h.element).startswith("Databases:"))
    set_paragraph_text(skills.element, "Databases: PostgreSQL, Redis, MongoDB")
    runs = own_runs(skills.element)
    first = runs[0]
    assert paragraph_text(skills.element) == "Databases: PostgreSQL, Redis, MongoDB"
    assert first.find(f"{W}rPr/{W}b") is not None and "Databases:" in "".join(t.text for t in first.iter(f"{W}t"))
    header = next(h for h in handles if "Northwind" in paragraph_text(h.element))
    set_paragraph_text(header.element, paragraph_text(header.element))  # no-op keeps tab + italics
    out = tmp_path / "x.docx"
    d.save(out)
    d2 = docx.Document(str(out))
    assert any(paragraph_text(h.element) == "Databases: PostgreSQL, Redis, MongoDB" for h in enumerate_paragraphs(d2))


# ---------------------------------------------------------------- completeness + identity
def test_completeness_detects_loss(sample_docx):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    texts = {b.id: b.text for b in r.blocks}
    assert validate_completeness(r, texts)["passed"]
    broken = dict(texts)
    hdr = r.experience[1].header_block_ids[0]
    broken[hdr] = "Software Engineer | Some Other Company\tJun 2018 – Feb 2021"
    assert not validate_completeness(r, broken)["passed"]
    broken = dict(texts)
    for bid in r.experience[2].bullet_block_ids:
        broken[bid] = ""
    assert not validate_completeness(r, broken)["passed"]
    summary_only = {b.id: (b.text if b.role in ("summary", "name") else "") for b in r.blocks}
    assert not validate_completeness(r, summary_only)["passed"]


def test_completeness_allows_name_contact_outside_text(sample_docx):
    # Name/contact rendered as a header image: nothing detected in text, so nothing can be lost.
    r = parse_resume(sample_docx, "docx", "resume.docx")
    r.name, r.contact.email, r.contact.phone = "", None, None
    texts = {b.id: b.text for b in r.blocks}
    assert validate_completeness(r, texts)["passed"]
    # ...but a detected name that disappears from the output still fails.
    r2 = parse_resume(sample_docx, "docx", "resume.docx")
    gone = {b.id: b.text.replace(r2.name, "") for b in r2.blocks}
    assert not validate_completeness(r2, gone)["passed"]


def test_identity_rejects_demo_data(sample_docx):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    texts = {b.id: b.text for b in r.blocks}
    ok = validate_identity(r, texts, r.full_text)
    assert ok["passed"]
    bad = validate_identity(r, texts, r.full_text.replace("Jordan Avery", "John Doe"))
    assert not bad["passed"]


# ---------------------------------------------------------------- AI response handling
def test_parse_json_text_variants():
    assert parse_json_text('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json_text('Sure! {"a": 2} done') == {"a": 2}
    with pytest.raises(AIResponseError):
        parse_json_text("not json")
    with pytest.raises(AIResponseError):
        parse_json_text('{"a": [1, 2')  # truncated


def test_ai_structure_rejects_unverified_claims(sample_docx, fake_ai):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    notes = ai_service.refine_structure(r)
    assert any("discarded" in n for n in notes)
    assert not any("Kubernetes" in c.claim for c in r.claims)
    assert any("Redis caching" in c.claim for c in r.claims)


def test_ai_jd_rejects_keywords_not_in_jd(sample_docx, sample_jd, fake_ai):
    r = parse_resume(sample_docx, "docx", "resume.docx")
    jd, evidence, notes = ai_service.analyze_jd(sample_jd, r, parse_jd_text(sample_jd))
    terms = {k.term for k in jd.keywords}
    assert "Haskell" not in terms and "REST APIs" in terms
    ev = {e.term: e for e in assess_keywords(jd, r.full_text, evidence)}
    assert ev["Kafka"].status == "missing"  # fake quote rejected
    assert ev["Microservices"].status == "partial" and ev["Microservices"].usable  # verified quote


def test_tailor_retries_incomplete_response(sample_docx, sample_jd):
    provider = FakeProvider(malformed_first=1)
    ai_service.set_provider(provider)
    try:
        r = parse_resume(sample_docx, "docx", "resume.docx")
        jd = parse_jd_text(sample_jd)
        props, notes = ai_service.tailor(r, jd, assess_keywords(jd, r.full_text), set(), "balanced")
        assert provider.calls.count("tailor") == 2
        assert props
    finally:
        ai_service.set_provider(None)


def test_ai_failure_falls_back(sample_docx, sample_jd):
    ai_service.set_provider(FakeProvider(fail_tasks={"tailor", "jd", "structure"}))
    try:
        r = parse_resume(sample_docx, "docx", "resume.docx")
        jd, _, notes = ai_service.analyze_jd(sample_jd, r, parse_jd_text(sample_jd))
        assert jd.source == "heuristic" and notes
        props, notes = ai_service.tailor(r, jd, assess_keywords(jd, r.full_text), set(), "balanced")
        assert props == [] and any("failed" in n for n in notes)
    finally:
        ai_service.set_provider(None)
