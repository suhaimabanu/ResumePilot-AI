"""Generate DEMO/TEST fixtures. These files are only used by the test suite.

The production workflow never reads from this directory; it only processes the user's upload.
Run: backend/.venv/Scripts/python fixtures/make_fixtures.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

HERE = Path(__file__).resolve().parent
ACCENT = RGBColor(0x2E, 0x4A, 0x7D)


def _photo(path: Path) -> None:
    img = Image.new("RGB", (300, 300), (210, 220, 235))
    d = ImageDraw.Draw(img)
    d.ellipse((90, 50, 210, 170), fill=(120, 130, 150))
    d.ellipse((50, 180, 250, 380), fill=(120, 130, 150))
    img.save(path)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = ACCENT
    # bottom border
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": "6", "space": "1", "color": "2E4A7D"}.items():
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _entry_header(doc: Document, left_bold: str, left_rest: str, dates: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_bold)
    r.bold = True
    p.add_run(left_rest)
    d = p.add_run("\t" + dates)
    d.italic = True


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)


def make_resume(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    photo = HERE / "_photo.png"
    _photo(photo)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    left.width, right.width = Inches(5.6), Inches(1.4)
    p = left.paragraphs[0]
    r = p.add_run("Jordan Avery")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = ACCENT
    p2 = left.add_paragraph()
    p2.add_run("Backend Software Engineer").font.size = Pt(12)
    p3 = left.add_paragraph()
    p3.add_run("jordan.avery@example.com | +1 555 010 2030 | linkedin.com/in/jordan-avery-demo | Austin, TX")
    right.paragraphs[0].add_run().add_picture(str(photo), width=Inches(1.2))

    _heading(doc, "Professional Summary")
    doc.add_paragraph(
        "Software engineer with 6 years of experience building backend services in Python. "
        "Experienced with Django, FastAPI and PostgreSQL, deploying services on AWS with Docker. "
        "Comfortable working in agile teams and mentoring junior developers."
    )

    _heading(doc, "Technical Skills")
    for label, items in [("Languages: ", "Python, SQL, JavaScript, Bash"),
                         ("Frameworks: ", "Django, FastAPI, Flask, Celery"),
                         ("Cloud & DevOps: ", "AWS (EC2, S3, Lambda), Docker, GitHub Actions, Linux"),
                         ("Databases: ", "Postgres, Redis, MongoDB")]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        lr = p.add_run(label)
        lr.bold = True
        p.add_run(items)

    _heading(doc, "Professional Experience")
    _entry_header(doc, "Senior Backend Engineer", " | Northwind Analytics", "Mar 2021 – Present")
    _bullet(doc, "Developed RESTful services in Python and FastAPI for the customer analytics platform.")
    _bullet(doc, "Reduced average API response time by 35% by adding Redis caching and query optimisation.")
    _bullet(doc, "Containerised 12 services with Docker and deployed them to AWS ECS using GitHub Actions pipelines.")
    _bullet(doc, "Mentored 3 junior developers through code reviews and pair programming.")
    _entry_header(doc, "Software Engineer", " | Bluebird Logistics", "Jun 2018 – Feb 2021")
    _bullet(doc, "Built backend applications in Django to manage shipment tracking for 200+ clients.")
    _bullet(doc, "Designed PostgreSQL schemas and wrote data migration scripts.")
    _bullet(doc, "Worked in two-week sprints with product and QA teams.")
    _entry_header(doc, "Junior Developer", " | Copperleaf Studio", "Jan 2017 – May 2018")
    _bullet(doc, "Maintained internal tools written in Python and Bash.")
    _bullet(doc, "Automated weekly reporting tasks, saving the operations team 5 hours per week.")

    _heading(doc, "Projects")
    _entry_header(doc, "Open Source Task Queue Dashboard", "", "2022")
    _bullet(doc, "Created a Flask dashboard to monitor Celery task queues, used by 40+ GitHub users.")
    _entry_header(doc, "Expense Tracker API", "", "2020")
    _bullet(doc, "Built an expense tracking API with FastAPI and MongoDB with JWT authentication.")

    _heading(doc, "Education")
    _entry_header(doc, "B.Sc. Computer Science", " | University of Texas at Austin", "2012 – 2016")

    _heading(doc, "Certifications")
    doc.add_paragraph("AWS Certified Developer – Associate (2022)")

    footer = sec.footer.paragraphs[0]
    footer.text = "Jordan Avery — Resume"
    doc.save(path)
    photo.unlink(missing_ok=True)


SAMPLE_JD = """Senior Python Backend Engineer — Example Corp

We are looking for a Senior Python Backend Engineer to design and build scalable REST APIs.

Responsibilities:
- Design, build and maintain REST APIs and microservices using Python and FastAPI
- Deploy and operate services on Amazon Web Services using Docker and Kubernetes
- Optimize PostgreSQL database performance and caching with Redis
- Build CI/CD pipelines and infrastructure as code with Terraform
- Mentor engineers and participate in code reviews
- Collaborate with product teams in an Agile environment

Requirements:
- 5+ years of professional experience in backend development with Python
- Strong experience with FastAPI or Django
- Experience with PostgreSQL and Redis
- Hands-on experience with AWS and Docker
- Experience with Kubernetes and Terraform
- Bachelor's degree in Computer Science or related field

Nice to have:
- AWS certification
- Experience with Kafka
- Experience with GraphQL
"""


if __name__ == "__main__":
    make_resume(HERE / "sample_resume.docx")
    (HERE / "sample_jd.txt").write_text(SAMPLE_JD, encoding="utf-8")
    print("fixtures written")
